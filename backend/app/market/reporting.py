"""Report generation: CSV, Excel, PDF tearsheets, and the daily digest.

The digest returns structured data rather than a formatted string so phase 07
can deliver the same content over Telegram without re-deriving any of it.
"""

from __future__ import annotations

import csv
import io
from dataclasses import dataclass

# matplotlib must pick a headless backend before pyplot is imported.
import matplotlib
import numpy as np

from app.market.analytics import (
    Panel,
    instrument_stats,
    movers,
    sector_performance,
)

matplotlib.use("Agg")
import matplotlib.pyplot as plt

INK = "#131822"
MUTED = "#6B7488"
ACCENT = "#2F4FC4"
GAIN = "#0E7A54"
LOSS = "#B33A33"
GRID = "#DBDFEA"


# --- CSV --------------------------------------------------------------------

def bars_csv(panel: Panel, symbol: str) -> str:
    j = panel.symbols.index(symbol)
    buf = io.StringIO()
    writer = csv.writer(buf, lineterminator="\n")
    writer.writerow(["date", "close", "volume", "daily_return_pct"])
    closes = panel.close[:, j]
    for t, day in enumerate(panel.days):
        ret = (closes[t] / closes[t - 1] - 1) * 100 if t else 0.0
        writer.writerow([day.date().isoformat(), f"{closes[t]:.4f}",
                         int(panel.volume[t, j]), f"{ret:.4f}"])
    return buf.getvalue()


def universe_csv(panel: Panel) -> str:
    buf = io.StringIO()
    writer = csv.writer(buf, lineterminator="\n")
    writer.writerow(["symbol", "sector", "last", "period_return_pct",
                     "annualised_vol_pct", "max_drawdown_pct", "avg_volume"])
    for j, symbol in enumerate(panel.symbols):
        s = instrument_stats(panel, symbol)
        writer.writerow([s["symbol"], panel.sectors[j] or "",
                         s["last"], s["period_return_pct"], s["annualised_vol_pct"],
                         s["max_drawdown_pct"], s["avg_volume"]])
    return buf.getvalue()


# --- Excel ------------------------------------------------------------------

def market_workbook(panel: Panel) -> bytes:
    """A three-sheet workbook: universe stats, sector performance, index history."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="2F4FC4")

    def style_header(ws, columns: list[str]) -> None:
        ws.append(columns)
        for i, _ in enumerate(columns, start=1):
            cell = ws.cell(row=1, column=i)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="left")
            ws.column_dimensions[get_column_letter(i)].width = max(14, len(columns[i - 1]) + 3)
        ws.freeze_panes = "A2"

    # --- universe ---
    ws = wb.active
    ws.title = "Universe"
    style_header(ws, ["Symbol", "Sector", "Last", "Return %", "Ann. vol %",
                      "Sharpe", "Max DD %", "Up days %", "Avg volume"])
    for symbol in panel.symbols:
        s = instrument_stats(panel, symbol)
        sector = panel.sectors[panel.symbols.index(symbol)]
        ws.append([symbol, sector, s["last"], s["period_return_pct"],
                   s["annualised_vol_pct"], s["sharpe"], s["max_drawdown_pct"],
                   s["up_day_pct"], s["avg_volume"]])

    # --- sectors ---
    ws2 = wb.create_sheet("Sectors")
    style_header(ws2, ["Sector", "Names", "1D %", "1W %", "1M %", "YTD %"])
    windows = ["1d", "1w", "1m", "ytd"]
    perf = {w: {r["sector"]: r["return_pct"] for r in sector_performance(panel, w)}
            for w in windows}
    counts = {r["sector"]: r["constituents"] for r in sector_performance(panel, "1d")}
    for sector in sorted(counts):
        ws2.append([sector, counts[sector], *[perf[w].get(sector) for w in windows]])

    # --- index ---
    ws3 = wb.create_sheet("Index")
    style_header(ws3, ["Date", "Index", "Daily return %", "Total volume"])
    caps = (panel.close * panel.shares).sum(axis=1)
    index = 1000.0 * caps / caps[0]
    for t, day in enumerate(panel.days):
        ret = (index[t] / index[t - 1] - 1) * 100 if t else 0.0
        ws3.append([day.date(), round(float(index[t]), 2), round(float(ret), 3),
                    int(panel.volume[t].sum())])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# --- charts for the PDF -----------------------------------------------------

def _price_chart(panel: Panel, symbol: str) -> bytes:
    j = panel.symbols.index(symbol)
    days = [d.date() for d in panel.days]
    closes = panel.close[:, j]

    fig, (ax, axv) = plt.subplots(
        2, 1, figsize=(7.2, 3.6), height_ratios=[3, 1], sharex=True,
        gridspec_kw={"hspace": 0.08},
    )
    ax.plot(days, closes, color=ACCENT, linewidth=1.1)
    ax.fill_between(days, closes, closes.min() * 0.97, color=ACCENT, alpha=0.07)
    ax.set_ylabel("Price", color=MUTED, fontsize=8)
    ax.grid(True, color=GRID, linewidth=0.5)
    ax.set_axisbelow(True)

    axv.bar(days, panel.volume[:, j], color=MUTED, alpha=0.5, width=1.0)
    axv.set_ylabel("Volume", color=MUTED, fontsize=8)
    axv.grid(True, color=GRID, linewidth=0.5)
    axv.set_axisbelow(True)

    for a in (ax, axv):
        a.tick_params(colors=MUTED, labelsize=7)
        for spine in a.spines.values():
            spine.set_color(GRID)

    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=170, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return buf.getvalue()


def _drawdown_chart(panel: Panel, symbol: str) -> bytes:
    j = panel.symbols.index(symbol)
    closes = panel.close[:, j]
    peak = np.maximum.accumulate(closes)
    dd = (closes / peak - 1.0) * 100

    fig, ax = plt.subplots(figsize=(7.2, 1.6))
    ax.fill_between([d.date() for d in panel.days], dd, 0, color=LOSS, alpha=0.28)
    ax.plot([d.date() for d in panel.days], dd, color=LOSS, linewidth=0.8)
    ax.set_ylabel("Drawdown %", color=MUTED, fontsize=8)
    ax.grid(True, color=GRID, linewidth=0.5)
    ax.set_axisbelow(True)
    ax.tick_params(colors=MUTED, labelsize=7)
    for spine in ax.spines.values():
        spine.set_color(GRID)

    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=170, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return buf.getvalue()


# --- PDF --------------------------------------------------------------------

def tearsheet_pdf(panel: Panel, symbol: str, *, name: str, sector: str | None,
                  news: list[dict] | None = None) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import LETTER
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import (
        Image,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    # The charts below plot the entire history, so the stats table must describe
    # the same window. Using the 252-day default here printed a max drawdown of
    # -29.6% directly above a chart showing -50%.
    stats = instrument_stats(panel, symbol, lookback=0)
    styles = getSampleStyleSheet()
    title = ParagraphStyle("t", parent=styles["Title"], fontSize=19, spaceAfter=2,
                           textColor=colors.HexColor(INK), alignment=0)
    sub = ParagraphStyle("s", parent=styles["Normal"], fontSize=9.5,
                         textColor=colors.HexColor(MUTED), spaceAfter=14)
    h2 = ParagraphStyle("h2", parent=styles["Heading2"], fontSize=11,
                        textColor=colors.HexColor(INK), spaceBefore=14, spaceAfter=6)
    small = ParagraphStyle("sm", parent=styles["Normal"], fontSize=8.5,
                           textColor=colors.HexColor(MUTED))

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=LETTER,
        leftMargin=0.7 * inch, rightMargin=0.7 * inch,
        topMargin=0.7 * inch, bottomMargin=0.7 * inch,
        title=f"{symbol} tearsheet",
    )

    story: list = [
        Paragraph(f"{symbol} &nbsp;&mdash;&nbsp; {name}", title),
        Paragraph(
            f"{sector or 'Unclassified'} &nbsp;·&nbsp; simulated instrument &nbsp;·&nbsp; "
            f"{stats['lookback_days']:,} sessions, "
            f"{panel.days[0].date():%d %b %Y} to {panel.days[-1].date():%d %b %Y}",
            sub,
        ),
    ]

    rows = [
        ["Last", f"{stats['last']:,.2f}",
         "Annualised return", f"{stats['annualised_return_pct']}%"],
        ["Period return", f"{stats['period_return_pct']}%",
         "Annualised vol", f"{stats['annualised_vol_pct']}%"],
        ["Max drawdown", f"{stats['max_drawdown_pct']}%",
         "Sharpe", str(stats["sharpe"])],
        ["Best day", f"{stats['best_day_pct']}%",
         "Worst day", f"{stats['worst_day_pct']}%"],
        ["Up days", f"{stats['up_day_pct']}%",
         "Avg volume", f"{stats['avg_volume']:,}"],
    ]
    table = Table(rows, colWidths=[1.5 * inch, 1.55 * inch, 1.65 * inch, 1.4 * inch])
    table.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (-1, -1), "Helvetica"),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("TEXTCOLOR", (0, 0), (0, -1), colors.HexColor(MUTED)),
        ("TEXTCOLOR", (2, 0), (2, -1), colors.HexColor(MUTED)),
        ("FONTNAME", (1, 0), (1, -1), "Helvetica-Bold"),
        ("FONTNAME", (3, 0), (3, -1), "Helvetica-Bold"),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("LINEBELOW", (0, 0), (-1, -2), 0.4, colors.HexColor(GRID)),
    ]))
    story += [
        table,
        Paragraph("Price and volume", h2),
        Image(io.BytesIO(_price_chart(panel, symbol)),
              width=6.9 * inch, height=3.45 * inch),
        Paragraph("Drawdown from peak", h2),
        Image(io.BytesIO(_drawdown_chart(panel, symbol)),
              width=6.9 * inch, height=1.55 * inch),
    ]

    if news:
        story.append(Paragraph("Recent headlines", h2))
        for item in news[:6]:
            score = item.get("sentiment", 0)
            tone = GAIN if score > 0.1 else LOSS if score < -0.1 else MUTED
            story.append(Paragraph(
                f'<font color="{tone}">&#9632;</font> &nbsp;'
                f'{item["published_at"][:10]} &mdash; {item["headline"]}',
                small,
            ))
            story.append(Spacer(1, 3))

    story += [
        Spacer(1, 16),
        Paragraph(
            "Every figure above covers the full period shown in the charts. Sharpe "
            "assumes a zero risk-free rate. Simulated data: this instrument does not "
            "exist and these figures describe no real security. Not investment advice.",
            small),
    ]

    doc.build(story)
    return buf.getvalue()


# --- daily digest -----------------------------------------------------------

@dataclass
class Digest:
    as_of: str
    index_value: float
    index_change_pct: float
    regime: str
    advancers: int
    decliners: int
    top_gainers: list[dict]
    top_losers: list[dict]
    best_sector: dict | None
    worst_sector: dict | None

    def to_text(self) -> str:
        """Plain text, ready for Telegram in phase 07."""
        arrow = "▲" if self.index_change_pct >= 0 else "▼"
        lines = [
            f"Market close — {self.as_of}",
            f"{arrow} Index {self.index_value:,.1f} "
            f"({self.index_change_pct:+.2f}%) · {self.regime}",
            f"Breadth {self.advancers} up / {self.decliners} down",
            "",
        ]
        if self.best_sector:
            lines.append(f"Best sector: {self.best_sector['sector']} "
                         f"{self.best_sector['return_pct']:+.2f}%")
        if self.worst_sector:
            lines.append(f"Worst sector: {self.worst_sector['sector']} "
                         f"{self.worst_sector['return_pct']:+.2f}%")
        lines.append("")
        lines.append("Gainers: " + ", ".join(
            f"{g['symbol']} {g['change_pct']:+.1f}%" for g in self.top_gainers[:3]))
        lines.append("Losers:  " + ", ".join(
            f"{lo['symbol']} {lo['change_pct']:+.1f}%" for lo in self.top_losers[:3]))
        return "\n".join(lines)


def build_digest(panel: Panel, snapshot: dict) -> Digest:
    sectors = sector_performance(panel, "1d")
    m = movers(panel, "1d", limit=5)
    return Digest(
        as_of=panel.days[-1].date().isoformat(),
        index_value=float(snapshot["index_value"]),
        index_change_pct=float(snapshot["index_return"]) * 100,
        regime=snapshot.get("regime") or "unknown",
        advancers=int(snapshot["advancers"]),
        decliners=int(snapshot["decliners"]),
        top_gainers=m["gainers"],
        top_losers=m["losers"],
        best_sector=sectors[0] if sectors else None,
        worst_sector=sectors[-1] if sectors else None,
    )


__all__ = [
    "Digest",
    "bars_csv",
    "build_digest",
    "market_workbook",
    "tearsheet_pdf",
    "universe_csv",
]
